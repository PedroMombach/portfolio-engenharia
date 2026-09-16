"""Leitura e validação das entradas: JSON da emissão, catálogo e itens."""

from __future__ import annotations

from collections import OrderedDict
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import re

from openpyxl import load_workbook

CATALOG_COLUMNS = (
    "COD_FAMILIA", "COD_ELEMENTO", "COD_CLIENTE_ELEMENTO", "MATERIAL", "TIPO",
    "TITULO_FAMILIA", "DIMENSÃO CARACTERÍSTICA", "UNIDADE", "REF_COMERCIAL",
    "DESCR_FAMILIA",
)
ITEM_COLUMNS = ("COD_ELEMENTO", "QUANTIDADE")

# Colunas que devem ser idênticas em todos os elementos de uma mesma família.
FAMILY_COLUMNS = ("MATERIAL", "TIPO", "TITULO_FAMILIA", "REF_COMERCIAL", "DESCR_FAMILIA")

# Campo do JSON -> comando LaTeX usado pelo template.
METADATA = {
    "titulo": "Titulo",
    "codigo_emitente": "CodigoEmitente",
    "codigo_cliente": "CodigoCliente",
    "revisao": "Rev",
    "cliente": "Cliente",
    "contratada": "Contratada",
    "disciplina": "Disciplina",
}
REVISION_COLUMNS = ("revisao", "data", "descricao", "elaborador", "verificador", "aprovador")

WINDOWS_RESERVED = (
    {"CON", "PRN", "AUX", "NUL"}
    | {f"COM{i}" for i in range(1, 10)}
    | {f"LPT{i}" for i in range(1, 10)}
)


class InputError(ValueError):
    """Entrada inválida, com mensagem destinada ao usuário."""


def text(value) -> str:
    return "" if value is None else str(value).strip()


# ---------------------------------------------------------------------------
# JSON da emissão
# ---------------------------------------------------------------------------

def _check_identifier(identifier: str) -> None:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", identifier):
        raise InputError("identificador: use até 80 letras ASCII, números, hífen ou sublinhado.")
    if identifier.upper() in WINDOWS_RESERVED:
        raise InputError("identificador reservado pelo Windows.")


def _check_revisions(revisions, current_revision: str) -> None:
    if not isinstance(revisions, list) or not revisions:
        raise InputError("historico_revisoes deve conter pelo menos uma revisão.")
    for index, revision in enumerate(revisions, 1):
        complete = isinstance(revision, dict) and all(
            isinstance(revision.get(key), str) and revision[key].strip()
            for key in REVISION_COLUMNS
        )
        if not complete:
            raise InputError(
                f"Histórico, linha {index}: preencha {', '.join(REVISION_COLUMNS)}."
            )
    if text(revisions[-1]["revisao"]) != current_revision:
        raise InputError("A última revisão do histórico deve coincidir com revisao.")


def load_emission(path) -> dict:
    """Lê o JSON da emissão e resolve os caminhos relativos a ele."""
    path = Path(path).resolve()
    try:
        config = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as exc:
        raise InputError(f"Não foi possível ler o JSON {path}: {exc}") from exc
    if not isinstance(config, dict):
        raise InputError("A configuração deve ser um objeto JSON.")

    for field in ("identificador", *METADATA, "catalogo", "figuras", "itens"):
        value = config.get(field)
        if not isinstance(value, str) or not value.strip():
            raise InputError(f"Campo obrigatório de texto: {field}")
        config[field] = value.strip()

    _check_identifier(config["identificador"])

    page = config.get("pagina_inicial", 1)
    if type(page) is not int or page < 1:
        raise InputError("pagina_inicial deve ser um inteiro positivo.")
    config["pagina_inicial"] = page

    _check_revisions(config.get("historico_revisoes"), config["revisao"])

    for key in ("catalogo", "figuras", "itens"):
        config[key] = (path.parent / config[key]).resolve()
    config["arquivo"] = path
    return config


# ---------------------------------------------------------------------------
# Planilhas
# ---------------------------------------------------------------------------

def read_sheet(path, sheet_name: str, columns) -> list[dict]:
    """Lê valores literais de uma aba; fórmulas não são aceitas como entrada."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise InputError(f"Não foi possível abrir {path}: {exc}") from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise InputError(f"{path}: aba '{sheet_name}' não encontrada.")
        rows = workbook[sheet_name].iter_rows()
        header = [text(cell.value) for cell in next(rows, [])]
        if any(header.count(column) != 1 for column in columns):
            raise InputError(
                f"{path}, aba {sheet_name}: cabeçalhos obrigatórios, "
                f"sem duplicação: {', '.join(columns)}"
            )
        positions = {column: header.index(column) for column in columns}

        records = []
        for number, row in enumerate(rows, 2):
            if all(cell.value is None for cell in row):
                continue
            if any(row[pos].data_type in {"f", "e"} for pos in positions.values()):
                raise InputError(
                    f"{path}, linha {number}: use valores literais, sem fórmulas ou erros Excel."
                )
            record = {column: row[pos].value for column, pos in positions.items()}
            record["_linha"] = number
            records.append(record)

        if not records:
            raise InputError(f"{path}, aba {sheet_name}: nenhum item informado.")
        return records
    finally:
        workbook.close()


def validate_catalog(records) -> list[dict]:
    """Garante códigos únicos e famílias consistentes."""
    elements, client_codes, families = set(), set(), {}
    result = []
    for source in records:
        row = {column: text(source.get(column)) for column in CATALOG_COLUMNS}
        where = f"Catálogo, linha {source.get('_linha', '?')}"

        for column in CATALOG_COLUMNS:
            if column != "REF_COMERCIAL" and not row[column]:
                raise InputError(f"{where}: {column} ausente.")

        code, family = row["COD_ELEMENTO"], row["COD_FAMILIA"]
        client_code = row["COD_CLIENTE_ELEMENTO"]
        if code in elements:
            raise InputError(f"{where}: COD_ELEMENTO duplicado: {code}")
        if client_code in client_codes:
            raise InputError(f"{where}: COD_CLIENTE_ELEMENTO duplicado: {client_code}")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]*", family):
            raise InputError(f"{where}: código de família inválido: {family}")

        shared = tuple(row[column] for column in FAMILY_COLUMNS)
        if families.setdefault(family, shared) != shared:
            raise InputError(f"{where}: dados inconsistentes na família {family}.")

        elements.add(code)
        client_codes.add(client_code)
        result.append(row)

    if not result:
        raise InputError("Catálogo vazio.")
    return result


def _parse_quantity(code: str, value) -> Decimal:
    if isinstance(value, (bool, str)) or value is None:
        raise InputError(f"{code}: QUANTIDADE deve ser uma célula numérica preenchida.")
    try:
        quantity = Decimal(str(value))
    except InvalidOperation as exc:
        raise InputError(f"{code}: QUANTIDADE inválida.") from exc
    if not quantity.is_finite() or quantity < 0:
        raise InputError(f"{code}: QUANTIDADE deve ser finita e não negativa.")
    return quantity


def validate_items(records) -> dict[str, Decimal]:
    """Converte a aba Itens em {código: quantidade}."""
    items = {}
    for row in records:
        code = text(row.get("COD_ELEMENTO"))
        if not code:
            raise InputError(f"Itens, linha {row.get('_linha', '?')}: COD_ELEMENTO ausente.")
        if code in items:
            raise InputError(f"COD_ELEMENTO duplicado na emissão: {code}")
        items[code] = _parse_quantity(code, row.get("QUANTIDADE"))
    if not items:
        raise InputError("A emissão precisa conter pelo menos um item.")
    return items


def select_materials(catalog, items) -> OrderedDict:
    """Agrupa os itens com quantidade positiva em material -> tipo -> família.

    Materiais mantêm a ordem do catálogo; tipos são ordenados alfabeticamente.
    Quantidade zero é aceita na entrada e omitida do documento.
    """
    unknown = set(items) - {row["COD_ELEMENTO"] for row in catalog}
    if unknown:
        raise InputError("Códigos não encontrados no catálogo: " + ", ".join(sorted(unknown)))

    materials = OrderedDict()
    for row in catalog:
        quantity = items.get(row["COD_ELEMENTO"])
        if not quantity:
            continue
        types = materials.setdefault(row["MATERIAL"], OrderedDict())
        families = types.setdefault(row["TIPO"], OrderedDict())
        families.setdefault(row["COD_FAMILIA"], []).append({**row, "QUANTIDADE": quantity})

    if not materials:
        raise InputError("A emissão não contém elementos com quantidade positiva.")

    def by_type(entry):
        return entry[0].casefold(), entry[0]

    return OrderedDict(
        (material, OrderedDict(sorted(types.items(), key=by_type)))
        for material, types in materials.items()
    )
