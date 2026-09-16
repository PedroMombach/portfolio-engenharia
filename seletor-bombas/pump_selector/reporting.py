"""Selection-report values for an Excel-authored, chart-bearing template."""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
from openpyxl.utils import get_column_letter

from .models import OperatingPoint, PumpCurve
from .selection import CandidateEvaluation, SelectionResult
from .workbook_xml import write_template_copy


def _curve_cells(
    start_col: int, curve: PumpCurve, speed_rpm: float,
    synchronous_speed_rpm: float, slip: float,
) -> dict[str, object]:
    """Values for a pump sheet's reserved 31-point block."""
    cells: dict[str, object] = {}
    for offset, value in enumerate(
        (speed_rpm, curve.frequency_hz, synchronous_speed_rpm, slip)
    ):
        cells[f"{get_column_letter(start_col + offset)}18"] = value
    for row, (q, h, eta, npsh) in enumerate(
        zip(curve.q_lps, curve.head_m, curve.efficiency, curve.npshr_m), 21
    ):
        values = (
            float(q), float(h), float(eta),
            None if math.isnan(float(npsh)) else float(npsh),
        )
        for offset, value in enumerate(values):
            cells[f"{get_column_letter(start_col + offset)}{row}"] = value
    return cells


def _operation_values(operation: OperatingPoint) -> tuple[object, ...]:
    npsh = operation.npsh
    return (
        npsh.suction_head_m,
        npsh.suction_loss_m,
        npsh.available_m,
        "Não informado" if npsh.required_m is None else npsh.required_m,
        npsh.available_to_required_ratio,
        npsh.absolute_margin_m,
    )


def _summary_cells(
    result: SelectionResult, evaluations: list[CandidateEvaluation],
    operations: list[OperatingPoint],
) -> dict[str, object]:
    request = result.request
    atmospheric = operations[0].npsh.atmospheric_m
    cells: dict[str, object] = {
        "C6": request.altitude_m,
        "C7": request.temperature_c,
        "C8": atmospheric,
        "C9": request.vapor_pressure_head_m,
    }
    for column, (evaluation, operation) in enumerate(zip(evaluations, operations), 3):
        letter = get_column_letter(column)
        metadata = evaluation.pump.metadata
        technical = (
            metadata.manufacturer,
            metadata.model,
            metadata.motor_power_kw,
            metadata.poles,
            "Não informado" if metadata.mass_kg is None else metadata.mass_kg,
            metadata.voltage_v,
            metadata.minimum_submergence_mm,
        )
        performance = (
            metadata.manufacturer,
            metadata.model,
            operation.q_lps,
            operation.head_m,
            operation.hydraulic_efficiency,
            operation.speed_rpm,
            operation.shaft_power_kw,
            metadata.motor_efficiency,
            operation.electrical_power_kw,
            operation.motor_margin,
        )
        cavitation = (
            metadata.manufacturer, metadata.model, *_operation_values(operation)
        )
        for first_row, values in ((13, technical), (23, performance), (36, cavitation)):
            for row, value in enumerate(values, first_row):
                cells[f"{letter}{row}"] = value
    return cells


def _minimum_operation(evaluation: CandidateEvaluation, result: SelectionResult) -> OperatingPoint:
    request = result.request
    return evaluation.pump.set_op(
        name="Minimum",
        altitude_m=request.altitude_m,
        q_lps=request.minimum_system.reference_flow_lps,
        head_m=request.minimum_system.reference_head_m,
        suction_loss_m=request.suction_loss_m,
        vapor_pressure_head_m=request.vapor_pressure_head_m,
    )


def _pump_cells(evaluation: CandidateEvaluation, minimum: OperatingPoint) -> dict[str, object]:
    pump = evaluation.pump
    maximum = evaluation.maximum_operation
    if maximum is None:
        raise ValueError(f"selected pump {pump.ID} has no maximum operating point")
    base_sync = pump.metadata.base_synchronous_speed_rpm
    cells = _curve_cells(
        2, pump.base_curve, pump.metadata.base_speed_rpm,
        base_sync, pump.metadata.base_slip,
    )
    for start_col, operation in ((7, maximum), (12, minimum)):
        cells.update(_curve_cells(
            start_col, pump.curve_at(operation.frequency_hz),
            operation.speed_rpm, operation.synchronous_speed_rpm, operation.slip,
        ))
    return cells


def _system_cells(
    result: SelectionResult, evaluations: list[CandidateEvaluation],
    minima: list[OperatingPoint],
) -> dict[str, object]:
    maximum_system = result.request.maximum_system
    minimum_system = result.request.minimum_system
    q_limits = [
        maximum_system.reference_flow_lps * 1.5,
        minimum_system.reference_flow_lps * 1.5,
    ]
    for evaluation, minimum in zip(evaluations, minima):
        maximum = evaluation.maximum_operation
        if maximum is None:
            continue
        q_limits.extend((
            float(evaluation.pump.curve_at(maximum.frequency_hz).q_lps[-1]),
            float(evaluation.pump.curve_at(minimum.frequency_hz).q_lps[-1]),
        ))
    cells: dict[str, object] = {}
    for row, q in enumerate(np.linspace(0, max(q_limits), 61), 2):
        cells[f"A{row}"] = float(q)
        cells[f"B{row}"] = float(minimum_system.head(float(q)))
        cells[f"C{row}"] = float(maximum_system.head(float(q)))
    return cells


def write_selection_report(
    result: SelectionResult, output_path: str | Path, template_path: str | Path,
) -> Path:
    """Populate both duty cases while preserving every template chart part."""
    selected = result.selected[:3]
    if not selected:
        raise ValueError("selection has no pumps to report")
    maxima = [evaluation.maximum_operation for evaluation in selected]
    if any(operation is None for operation in maxima):
        raise ValueError("selected pump has no maximum operating point")
    minima = [_minimum_operation(evaluation, result) for evaluation in selected]
    updates = {
        "CurvaMax": _summary_cells(result, selected, maxima),
        "CurvaMin": _summary_cells(result, selected, minima),
        "CurvasSistema": _system_cells(result, selected, minima),
    }
    for index in range(1, 4):
        updates[f"Bomba {index}"] = (
            _pump_cells(selected[index - 1], minima[index - 1])
            if index <= len(selected) else {}
        )
    hidden = {f"Bomba {index}" for index in range(len(selected) + 1, 4)}
    return write_template_copy(
        template_path, output_path, updates,
        f"Seleção de bombas — {result.request.name}", hidden,
    )
