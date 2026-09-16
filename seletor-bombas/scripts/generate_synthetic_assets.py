"""Generate reproducible, non-proprietary XLSX assets for the public project."""

from __future__ import annotations

from pathlib import Path
import math
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pump_selector.database import CURVE_COLUMNS, METADATA_COLUMNS, PumpDatabase
from pump_selector.models import SystemCurve
from pump_selector.selection import SelectionCriteria, SelectionRequest, evaluate_pump


BLUE = "0077B6"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EEF8"
FONT = "Arial"


def _choose_motor_power(required_kw: float) -> float:
    standard = [0.75, 1.1, 1.5, 2.2, 3.0, 4.0, 5.5, 7.5, 11.0, 15.0, 18.5, 22.0, 30.0, 37.0, 45.0, 55.0, 75.0]
    for value in standard:
        if value >= required_kw * 1.18:
            return value
    return math.ceil(required_kw * 1.20 / 5) * 5.0


def build_synthetic_frames() -> tuple[pd.DataFrame, pd.DataFrame]:
    manufacturers = [
        ("AquaNova", "AN", 0.96, 1.05, 0.000, -10),
        ("HidroVale", "HV", 1.00, 1.00, 0.015, 0),
        ("Fluxora", "FX", 1.04, 0.96, 0.025, 10),
    ]
    q_bep_values = [4.0, 7.0, 12.0, 20.0, 35.0, 60.0, 95.0]
    head_bep_values = [32.0, 30.0, 28.0, 25.0, 22.0, 18.0, 14.0]
    eta_values = [0.55, 0.60, 0.65, 0.70, 0.74, 0.77, 0.80]
    metadata_rows: list[dict[str, object]] = []
    curve_rows: list[dict[str, object]] = []

    for manufacturer, prefix, q_scale, h_scale, eta_delta, speed_delta in manufacturers:
        for index, (q_base, h_base, eta_base) in enumerate(
            zip(q_bep_values, head_bep_values, eta_values), 1
        ):
            pump_id = f"{prefix}-{index:03d}"
            q_bep = q_base * q_scale
            h_bep = h_base * h_scale
            eta_bep = min(0.86, eta_base + eta_delta)
            poles = 2 if index <= 2 else 4
            synchronous = 3600 if poles == 2 else 1800
            base_speed = synchronous - (95 if poles == 2 else 45) + speed_delta
            motor_efficiency = min(0.95, 0.80 + index * 0.018 + eta_delta / 2)
            hydraulic_power = 1e-3 * q_bep * 0.9978 * 9.80665 * h_bep / eta_bep
            motor_power = _choose_motor_power(hydraulic_power)
            mass = round(22 + 8.5 * motor_power + index * 2.5, 1)
            minimum_submergence = round(210 + q_bep * 2.1 + index * 5, 1)
            impeller_axis = round(145 + q_bep * 0.7, 1)

            x = np.linspace(0.0, 1.5, 31)
            q = q_bep * x
            head = h_bep * (1.35 - 0.35 * x**2)
            eta = eta_bep * np.clip(1 - (x - 1) ** 2, 0, None)
            npsh_missing = (prefix == "AN" and index == 2) or (prefix == "FX" and index == 6)
            npsh = np.full_like(x, np.nan) if npsh_missing else 0.35 + 0.75 * x + 1.10 * x**2

            bep_index = int(np.argmax(eta))
            metadata_rows.append(
                {
                    "ID": pump_id,
                    "manufacturer": manufacturer,
                    "model": f"{prefix}-S{index:02d}",
                    "poles": poles,
                    "baseSpeedRPM": float(base_speed),
                    "baseFreq": 60.0,
                    "motorEfficiency": round(motor_efficiency, 4),
                    "motorPowerKW": motor_power,
                    "minSubmergenceMM": minimum_submergence,
                    "impellerAxisMM": impeller_axis,
                    "voltageV": 380,
                    "massKG": mass,
                    "Qbep": float(q[bep_index]),
                    "Hbep": float(head[bep_index]),
                    "Eta1bep": float(eta[bep_index]),
                    "NPSHbep": None if npsh_missing else float(npsh[bep_index]),
                }
            )
            for q_value, head_value, eta_value, npsh_value in zip(q, head, eta, npsh):
                curve_rows.append(
                    {
                        "ID": pump_id,
                        "Q": round(float(q_value), 6),
                        "H": round(float(head_value), 6),
                        "Eta1": round(float(eta_value), 9),
                        "NPSH": None if math.isnan(float(npsh_value)) else round(float(npsh_value), 6),
                    }
                )
    return (
        pd.DataFrame(metadata_rows).loc[:, METADATA_COLUMNS],
        pd.DataFrame(curve_rows).loc[:, CURVE_COLUMNS],
    )


def build_abacus(metadata: pd.DataFrame, curves: pd.DataFrame) -> dict[str, pd.DataFrame]:
    base = PumpDatabase(metadata, curves, {})
    flows = [3, 5, 8, 12, 20, 35, 60, 90]
    heads = [40, 35, 30, 25, 20, 15, 10]
    criteria = SelectionCriteria(missing_npsh_policy="allow_with_warning")
    result: dict[str, pd.DataFrame] = {}
    for manufacturer in base.manufacturers:
        pump_ids = base.metadata.loc[base.metadata["manufacturer"] == manufacturer, "ID"].tolist()
        rows = []
        for head in heads:
            row: dict[object, object] = {"H\\Q": head}
            for flow in flows:
                request = SelectionRequest(
                    name="Abacus cell",
                    maximum_system=SystemCurve("Maximum", flow, head, 0.0),
                    minimum_system=SystemCurve("Minimum", flow, head, 0.0),
                    altitude_m=100.0,
                    top_n=1,
                )
                evaluations = [evaluate_pump(base.get_pump(pump_id), request, criteria) for pump_id in pump_ids]
                feasible = sorted(
                    (item for item in evaluations if item.feasible),
                    key=lambda item: (item.score, item.pump_id),
                )
                row[flow] = feasible[0].pump_id if feasible else None
            rows.append(row)
        result[manufacturer] = pd.DataFrame(rows, columns=["H\\Q", *flows])
    return result


def _write_dataframe(ws, frame: pd.DataFrame, table_name: str) -> None:
    # Excel tables require textual headers.  Numeric abacus axes are restored
    # as floats by the database loader when looking up the nearest duty point.
    ws.append([str(column) for column in frame.columns])
    for record in frame.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in record])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    table = Table(displayName=table_name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    for column in ws.columns:
        width = min(28, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width


def write_database(path: Path, metadata: pd.DataFrame, curves: pd.DataFrame, abacus: dict[str, pd.DataFrame]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    metadata_ws = workbook.create_sheet("Metadata")
    _write_dataframe(metadata_ws, metadata, "PumpMetadata")
    for column in (5, 6, 7, 8, 13, 14, 15, 16):
        for cell in metadata_ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "0.000"

    curves_ws = workbook.create_sheet("Curves")
    _write_dataframe(curves_ws, curves, "PumpCurves")
    for column in (2, 3, 5):
        for cell in curves_ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "0.000"
    for item in curves_ws.iter_cols(min_col=4, max_col=4, min_row=2):
        for cell in item:
            cell.number_format = "0.0%"

    for manufacturer, grid in abacus.items():
        ws = workbook.create_sheet(f"Abacus_{manufacturer}")
        _write_dataframe(ws, grid, f"Abacus{manufacturer}")
        ws.sheet_properties.tabColor = "5B9BD5"
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                if cell.value is None:
                    cell.fill = PatternFill("solid", fgColor="F2F2F2")
                else:
                    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
                cell.alignment = Alignment(horizontal="center", vertical="center")

    readme = workbook.create_sheet("ReadMe")
    readme.sheet_view.showGridLines = False
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 88
    notes = [
        ("Arquivo", "Banco sintético para demonstração pública; não representa produtos comerciais."),
        ("Metadata", "Um registro por ID único. baseFreq está em Hz; Q em L/s; alturas em m.c.a."),
        ("Curves", "Curvas sintéticas H-Q, eficiência e NPSHr em formato longo."),
        ("Abacus_*", "Ábacos sintéticos por fabricante. Cada célula contém um ID completo e independente."),
        ("NPSH vazio", "Significa não informado. O software não converte ausência em zero."),
        ("BEP", "Q, H, Eta1 e NPSH correspondem ao ponto amostrado de máxima eficiência."),
    ]
    for row, (label, text) in enumerate(notes, 1):
        readme.cell(row=row, column=1, value=label).font = Font(name=FONT, bold=True, color=WHITE)
        readme.cell(row=row, column=1).fill = PatternFill("solid", fgColor=BLUE)
        readme.cell(row=row, column=2, value=text).font = Font(name=FONT)
        readme.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    workbook.properties.creator = "Pump Selector contributors"
    workbook.properties.lastModifiedBy = "Pump Selector contributors"
    workbook.properties.title = "Synthetic pump database"
    workbook.save(path)


def write_example_inputs(path: Path) -> None:
    columns = [
        "Name",
        "Q_Lps",
        "StaticHeadMin_m",
        "StaticHeadMax_m",
        "HeadMin_m",
        "HeadMax_m",
        "Altitude_m",
        "Temperature_C",
        "SuctionLoss_m",
        "VaporHead_m",
        "Mode",
        "ManualIDs",
    ]
    rows = [
        ["Estação sintética — otimização", 12, 4.0, 5.0, 17.0, 24.0, 120, 20, 0.35, 0.26, "optimize", ""],
        ["Estação sintética — ábaco", 20, 3.0, 4.5, 15.0, 22.0, 80, 20, 0.25, 0.26, "abacus", ""],
        ["Estação sintética — manual", 7, 2.0, 3.0, 14.0, 20.0, 50, 20, 0.20, 0.26, "manual", "AN-002;HV-002;FX-002"],
    ]
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Cases"
    _write_dataframe(ws, pd.DataFrame(rows, columns=columns), "SelectionCases")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["L"].width = 34
    workbook.properties.creator = "Pump Selector contributors"
    workbook.properties.lastModifiedBy = "Pump Selector contributors"
    workbook.properties.title = "Synthetic pump selection inputs"
    workbook.save(path)


def main() -> None:
    metadata, curves = build_synthetic_frames()
    abacus = build_abacus(metadata, curves)
    database_path = PROJECT_ROOT / "DB" / "PumpDatabase.xlsx"
    template_path = PROJECT_ROOT / "TemplateSaida.xlsx"
    input_path = PROJECT_ROOT / "ExampleInputs.xlsx"
    database_path.parent.mkdir(parents=True, exist_ok=True)
    write_database(database_path, metadata, curves, abacus)
    # The report template is a user-maintained artifact.  Regenerating the
    # database must not discard manual edits or chart definitions.
    if not template_path.exists():
        raise FileNotFoundError(f"required report template is missing: {template_path}")
    write_example_inputs(input_path)
    print(f"created {database_path}")
    print(f"preserved {template_path}")
    print(f"created {input_path}")


if __name__ == "__main__":
    main()
