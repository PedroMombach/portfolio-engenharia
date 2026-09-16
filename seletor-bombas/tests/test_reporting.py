from pathlib import Path
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from pump_selector import PumpDatabase, load_cases, select_pumps, write_selection_report


ROOT = Path(__file__).resolve().parents[1]


class ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.database = PumpDatabase.from_xlsx(ROOT / "DB" / "PumpDatabase.xlsx")
        cls.cases = load_cases(ROOT / "ExampleInputs.xlsx")
        cls.output_directory = ROOT / ".test-output"
        cls.output_directory.mkdir(exist_ok=True)

    @classmethod
    def tearDownClass(cls):
        for name in ("report.xlsx", "missing-npsh.xlsx"):
            path = cls.output_directory / name
            if path.exists():
                path.unlink()
        cls.output_directory.rmdir()

    def test_report_populates_both_cases_system_curves_and_template_charts(self):
        case = self.cases[0]
        result = select_pumps(self.database, case.request(), case.mode)
        output = self.output_directory / "report.xlsx"
        write_selection_report(result, output, ROOT / "TemplateSaida.xlsx")
        template = load_workbook(ROOT / "TemplateSaida.xlsx")
        workbook = load_workbook(output, data_only=False)
        self.assertEqual(workbook.sheetnames, template.sheetnames)
        self.assertEqual(workbook["CurvaMax"]["C13"].value, result.selected[0].pump.metadata.manufacturer)
        self.assertEqual(workbook["CurvaMin"]["C13"].value, result.selected[0].pump.metadata.manufacturer)
        self.assertEqual(workbook["CurvaMax"]["C26"].value, result.selected[0].maximum_operation.head_m)
        self.assertAlmostEqual(workbook["CurvaMin"]["C26"].value, case.head_min_m)
        system = workbook["CurvasSistema"]
        self.assertEqual(system.sheet_state, "visible")
        self.assertEqual(system["A2"].value, 0)
        self.assertEqual(system["B2"].value, case.static_head_min_m)
        self.assertEqual(system["C2"].value, case.static_head_max_m)
        for row in (2, 31, 62):
            q = system.cell(row, 1).value
            self.assertAlmostEqual(system.cell(row, 2).value, case.request().minimum_system.head(q))
            self.assertAlmostEqual(system.cell(row, 3).value, case.request().maximum_system.head(q))
        self.assertEqual(system.max_column, 3)
        self.assertEqual(system["B2"].value, case.static_head_min_m)
        self.assertEqual(system["C2"].value, case.static_head_max_m)
        for sheet in template.sheetnames:
            self.assertEqual(len(workbook[sheet]._charts), len(template[sheet]._charts))
        with ZipFile(ROOT / "TemplateSaida.xlsx") as original, ZipFile(output) as generated:
            protected = [
                name for name in original.namelist()
                if name.startswith(("xl/charts/", "xl/drawings/"))
                or name in ("xl/styles.xml", "xl/theme/theme1.xml")
            ]
            self.assertTrue(protected)
            self.assertTrue(all(generated.read(name) == original.read(name) for name in protected))
        for sheet, cell in (("CurvaMax", "C6"), ("CurvaMin", "C26"),
                            ("Bomba 1", "B21"), ("CurvasSistema", "B2")):
            self.assertEqual(workbook[sheet][cell].number_format, template[sheet][cell].number_format)
        self.assertNotIn("ChartData", workbook.sheetnames)
        template.close()
        workbook.close()

    def test_missing_npsh_is_written_as_not_reported(self):
        case = self.cases[2]
        result = select_pumps(
            self.database,
            case.request(),
            case.mode,
            manual_ids=case.manual_ids,
        )
        output = self.output_directory / "missing-npsh.xlsx"
        write_selection_report(result, output, ROOT / "TemplateSaida.xlsx")
        workbook = load_workbook(output, data_only=True)
        self.assertEqual(workbook["CurvaMax"]["C41"].value, "Não informado")
        self.assertIsNone(workbook["CurvaMax"]["C42"].value)
        self.assertIsNone(workbook["CurvaMax"]["C43"].value)
        self.assertEqual(workbook["CurvaMin"]["C41"].value, "Não informado")
        self.assertIsNone(workbook["CurvaMin"]["C42"].value)
        self.assertIsNone(workbook["CurvaMin"]["C43"].value)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
